#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

import logging
import re
import struct
import sys
from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side

from rag.nlp import find_codec
from rag.utils.lazy_image import LazyImage

# copied from `/openpyxl/cell/cell.py`
ILLEGAL_CHARACTERS_RE = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")


def _biff8_cell_borders(stream):
    """Read explicit cell XF border geometry without evaluating XLS formulas.

    Calamine remains the authority for values and merges. Reject incomplete
    metadata as a unit so a truncated stream cannot prove a partial grid.
    Sheet slots include non-worksheets, matching BOUNDSHEET display order.
    """
    def records(offset):
        while offset < len(stream):
            if offset + 4 > len(stream):
                raise ValueError("truncated BIFF record header")
            kind, size = struct.unpack_from("<HH", stream, offset)
            end = offset + 4 + size
            if end > len(stream):
                raise ValueError("truncated BIFF record")
            yield kind, stream[offset + 4:end]
            offset = end
            if kind == 0x000A:
                return
        raise ValueError("missing BIFF EOF")

    def checked_records(offset, substream_type):
        iterator = records(offset)
        kind, payload = next(iterator, (None, b""))
        if kind != 0x0809 or len(payload) < 4 or struct.unpack_from("<HH", payload) != (0x0600, substream_type):
            raise ValueError("expected BIFF8 substream")
        return iterator

    formats, sheets = [], []
    for kind, payload in checked_records(0, 5):
        if kind == 0x00E0:
            if len(payload) != 20:
                raise ValueError("invalid BIFF8 XF")
            bits = struct.unpack_from("<I", payload, 10)[0]
            sides = tuple((bits >> shift) & 15 for shift in (0, 4, 8, 12))
            if any(side > 13 for side in sides):
                raise ValueError("invalid BIFF border style")
            formats.append(sides)
        elif kind == 0x0085:
            if len(payload) < 8:
                raise ValueError("invalid BIFF BOUNDSHEET")
            sheets.append((struct.unpack_from("<I", payload)[0], payload[5]))
    if not sheets:
        raise ValueError("missing BIFF sheets")
    result = []
    minimum_sizes = {0x0201: 6, 0x0203: 14, 0x0204: 9, 0x0205: 8, 0x00FD: 10, 0x027E: 10, 0x0006: 22}
    for offset, sheet_type in sheets:
        cells = {}
        result.append(cells)
        if sheet_type != 0:
            continue

        def add(row, column, xf):
            if column > 255 or xf >= len(formats):
                raise ValueError("invalid BIFF cell format reference")
            coordinate = (row + 1, column + 1)
            if coordinate in cells:
                raise ValueError("duplicate BIFF cell format")
            cells[coordinate] = formats[xf]

        for kind, payload in checked_records(offset, 16):
            if kind in minimum_sizes:
                if len(payload) < minimum_sizes[kind]:
                    raise ValueError("invalid BIFF cell record")
                add(*struct.unpack_from("<HHH", payload))
            elif kind in (0x00BE, 0x00BD):
                if len(payload) < 8:
                    raise ValueError("invalid BIFF multi-cell record")
                row, first = struct.unpack_from("<HH", payload)
                last = struct.unpack_from("<H", payload, len(payload) - 2)[0]
                stride = 2 if kind == 0x00BE else 6
                if last < first or len(payload) != 6 + (last - first + 1) * stride:
                    raise ValueError("invalid BIFF multi-cell extent")
                for index, column in enumerate(range(first, last + 1)):
                    add(row, column, struct.unpack_from("<H", payload, 4 + index * stride)[0])
    return result


class RAGFlowExcelParser:
    @staticmethod
    def _clean_cell_value(value):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub(" ", value)
            return value if value.strip() else None
        try:
            missing = pd.isna(value)
            if not hasattr(missing, "__len__") and bool(missing):
                return None
        except (TypeError, ValueError):
            pass
        return value

    @staticmethod
    def _load_calamine_to_workbook(file_like_object):
        from python_calamine import CalamineWorkbook

        file_like_object.seek(0)
        source = CalamineWorkbook.from_filelike(file_like_object)
        workbook = Workbook()
        workbook.remove(workbook.active)

        borders = []
        file_like_object.seek(0)
        if file_like_object.read(8) == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            import olefile

            file_like_object.seek(0)
            try:
                with olefile.OleFileIO(file_like_object) as ole:
                    stream_name = "Workbook" if ole.exists("Workbook") else "Book"
                    borders = _biff8_cell_borders(ole.openstream(stream_name).read())
                if len(borders) != len(source.sheet_names):
                    raise ValueError("BIFF/Calamine sheet count mismatch")
            except (ValueError, OSError, struct.error):
                # Missing visual evidence must never invent a grid. Values and
                # merges can still be parsed by Calamine without this channel.
                logging.warning("XLS border metadata unavailable")
                borders = []
        border_styles = (None, "thin", "medium", "dashed", "dotted", "thick", "double", "hair", "mediumDashed", "dashDot", "mediumDashDot", "dashDotDot", "mediumDashDotDot", "slantDashDot")
        border_cache = {}

        for sheet_index, sheet_name in enumerate(source.sheet_names):
            source_sheet = source.get_sheet_by_name(sheet_name)
            worksheet = workbook.create_sheet(title=sheet_name)
            sheet_start = getattr(source_sheet, "start", None)
            if sheet_start is None:
                continue

            for row_index, row in enumerate(source_sheet.iter_rows(), start=1):
                for column_index, value in enumerate(row, start=1):
                    cleaned = RAGFlowExcelParser._clean_cell_value(value)
                    if cleaned is not None:
                        worksheet.cell(row=row_index, column=column_index, value=cleaned)

            if borders:
                source_ranges = getattr(source_sheet, "merged_cell_ranges", [])
                max_row = max([worksheet.max_row, *(end[0] + 1 for _start, end in source_ranges)])
                max_column = max([worksheet.max_column, *(end[1] + 1 for _start, end in source_ranges)])
                for (row, column), sides in borders[sheet_index].items():
                    if row > max_row or column > max_column or not any(sides):
                        continue
                    if sides not in border_cache:
                        border_cache[sides] = Border(**{
                            name: Side(style=border_styles[style])
                            for name, style in zip(("left", "right", "top", "bottom"), sides)
                        })
                    worksheet.cell(row, column).border = border_cache[sides]

            for start, end in getattr(source_sheet, "merged_cell_ranges", []):
                min_row, min_col = start
                max_row, max_col = end
                if min_row != max_row or min_col != max_col:
                    worksheet.merge_cells(
                        start_row=min_row + 1,
                        start_column=min_col + 1,
                        end_row=max_row + 1,
                        end_column=max_col + 1,
                    )

        return workbook

    @staticmethod
    def _load_excel_to_workbook(file_like_object):
        if isinstance(file_like_object, bytes):
            file_like_object = BytesIO(file_like_object)

        # Read first 4 bytes to determine file type
        file_like_object.seek(0)
        file_head = file_like_object.read(4)
        file_like_object.seek(0)

        if not (file_head.startswith(b"PK\x03\x04") or file_head.startswith(b"\xd0\xcf\x11\xe0")):
            logging.info("Not an Excel file, converting CSV to Excel Workbook")

            try:
                file_like_object.seek(0)
                df = pd.read_csv(file_like_object, on_bad_lines="skip")
                return RAGFlowExcelParser._dataframe_to_workbook(df)

            except Exception as e_csv:
                raise Exception(f"Failed to parse CSV and convert to Excel Workbook: {e_csv}")

        try:
            return load_workbook(file_like_object, data_only=True)
        except Exception as e:
            logging.info(f"openpyxl load error: {e}, try calamine instead")
            try:
                try:
                    return RAGFlowExcelParser._load_calamine_to_workbook(file_like_object)
                except Exception as ex:
                    logging.info(f"calamine workbook load error: {ex}, try pandas instead")
                    file_like_object.seek(0)
                    dfs = pd.read_excel(file_like_object, sheet_name=None, header=None, engine="calamine")
                    return RAGFlowExcelParser._dataframes_to_workbook(dfs, include_headers=False)
            except Exception as e_pandas:
                raise Exception(f"pandas.read_excel error: {e_pandas}, original openpyxl error: {e}")

    @staticmethod
    def _clean_dataframe(df: pd.DataFrame):
        return df.apply(lambda col: col.map(RAGFlowExcelParser._clean_cell_value))

    @staticmethod
    def _fill_worksheet_from_dataframe(ws, df: pd.DataFrame, include_headers=True):
        data_start_row = 1
        if include_headers:
            for col_num, column_name in enumerate(df.columns, 1):
                ws.cell(row=1, column=col_num, value=column_name)
            data_start_row = 2
        for row_num, row in enumerate(df.values, data_start_row):
            for col_num, value in enumerate(row, 1):
                cleaned = RAGFlowExcelParser._clean_cell_value(value)
                if cleaned is not None:
                    ws.cell(row=row_num, column=col_num, value=cleaned)

    @staticmethod
    def _dataframe_to_workbook(df):
        if isinstance(df, dict) and len(df) > 1:
            return RAGFlowExcelParser._dataframes_to_workbook(df)

        df = RAGFlowExcelParser._clean_dataframe(df)
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        RAGFlowExcelParser._fill_worksheet_from_dataframe(ws, df)
        return wb

    @staticmethod
    def _dataframes_to_workbook(dfs: dict, include_headers=True):
        wb = Workbook()
        default_sheet = wb.active
        wb.remove(default_sheet)

        for sheet_name, df in dfs.items():
            df = RAGFlowExcelParser._clean_dataframe(df)
            ws = wb.create_sheet(title=sheet_name)
            RAGFlowExcelParser._fill_worksheet_from_dataframe(ws, df, include_headers=include_headers)
        return wb

    @staticmethod
    def _extract_images_from_worksheet(ws, sheetname=None):
        """
        Extract images from a worksheet and enrich them with vision-based descriptions.

        Returns: List[dict]
        """
        images = getattr(ws, "_images", [])
        if not images:
            return []

        raw_items = []

        for img in images:
            try:
                img_bytes = img._data()
                lazy_img = LazyImage([img_bytes])

                anchor = img.anchor
                if hasattr(anchor, "_from") and hasattr(anchor, "_to"):
                    r1, c1 = anchor._from.row + 1, anchor._from.col + 1
                    r2, c2 = anchor._to.row + 1, anchor._to.col + 1
                    if r1 == r2 and c1 == c2:
                        span = "single_cell"
                    else:
                        span = "multi_cell"
                else:
                    r1, c1 = anchor._from.row + 1, anchor._from.col + 1
                    r2, c2 = r1, c1
                    span = "single_cell"

                item = {
                    "sheet": sheetname or ws.title,
                    "image": lazy_img,
                    "image_description": "",
                    "row_from": r1,
                    "col_from": c1,
                    "row_to": r2,
                    "col_to": c2,
                    "span_type": span,
                }
                raw_items.append(item)
            except Exception:
                continue
        return raw_items

    @staticmethod
    def _get_actual_row_count(ws):
        max_row = ws.max_row
        if not max_row:
            return 0
        if max_row <= 10000:
            return max_row

        max_col = min(ws.max_column or 1, 50)

        def row_has_data(row_idx):
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None and str(cell.value).strip():
                    return True
            return False

        if not any(row_has_data(i) for i in range(1, min(101, max_row + 1))):
            return 0

        left, right = 1, max_row
        last_data_row = 1

        while left <= right:
            mid = (left + right) // 2
            found = False
            for r in range(mid, min(mid + 10, max_row + 1)):
                if row_has_data(r):
                    found = True
                    last_data_row = max(last_data_row, r)
                    break
            if found:
                left = mid + 1
            else:
                right = mid - 1

        for r in range(last_data_row, min(last_data_row + 500, max_row + 1)):
            if row_has_data(r):
                last_data_row = r

        return last_data_row

    @staticmethod
    def _get_rows_limited(ws):
        actual_rows = RAGFlowExcelParser._get_actual_row_count(ws)
        if actual_rows == 0:
            return []
        return list(ws.iter_rows(min_row=1, max_row=actual_rows))

    def html(self, fnm, chunk_rows=256):
        from html import escape

        file_like_object = BytesIO(fnm) if not isinstance(fnm, str) else fnm
        wb = RAGFlowExcelParser._load_excel_to_workbook(file_like_object)
        tb_chunks = []

        def _fmt(v):
            if v is None:
                return ""
            return str(v).strip()

        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            try:
                rows = RAGFlowExcelParser._get_rows_limited(ws)
            except Exception as e:
                logging.warning(f"Skip sheet '{sheetname}' due to rows access error: {e}")
                continue

            if not rows:
                continue

            tb_rows_0 = "<tr>"
            for t in list(rows[0]):
                tb_rows_0 += f"<th>{escape(_fmt(t.value))}</th>"
            tb_rows_0 += "</tr>"

            # rows[0] is the header; split the remaining data rows into
            # ceil(n_data / chunk_rows) chunks. Using +1 here over-counts by one
            # when the data-row count is an exact multiple of chunk_rows and emits
            # a spurious header-only chunk.
            n_data_rows = len(rows) - 1
            for chunk_i in range((n_data_rows + chunk_rows - 1) // chunk_rows):
                tb = ""
                tb += f"<table><caption>{sheetname}</caption>"
                tb += tb_rows_0
                for r in list(rows[1 + chunk_i * chunk_rows : min(1 + (chunk_i + 1) * chunk_rows, len(rows))]):
                    tb += "<tr>"
                    for i, c in enumerate(r):
                        if c.value is None:
                            tb += "<td></td>"
                        else:
                            tb += f"<td>{escape(_fmt(c.value))}</td>"
                    tb += "</tr>"
                tb += "</table>\n"
                tb_chunks.append(tb)

        return tb_chunks

    def markdown(self, fnm):
        import pandas as pd

        file_like_object = BytesIO(fnm) if not isinstance(fnm, str) else fnm
        try:
            file_like_object.seek(0)
            df = pd.read_excel(file_like_object)
        except Exception as e:
            logging.warning(f"Parse spreadsheet error: {e}, trying to interpret as CSV file")
            file_like_object.seek(0)
            df = pd.read_csv(file_like_object, on_bad_lines="skip")
        df = df.replace(r"^\s*$", "", regex=True)
        return df.to_markdown(index=False)

    def __call__(self, fnm):
        file_like_object = BytesIO(fnm) if not isinstance(fnm, str) else fnm
        wb = RAGFlowExcelParser._load_excel_to_workbook(file_like_object)

        res = []
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            try:
                rows = RAGFlowExcelParser._get_rows_limited(ws)
            except Exception as e:
                logging.warning(f"Skip sheet '{sheetname}' due to rows access error: {e}")
                continue
            if not rows:
                continue
            ti = list(rows[0])
            for r in list(rows[1:]):
                fields = []
                for i, c in enumerate(r):
                    if c.value is None or str(c.value).strip() == "":
                        continue
                    t = str(ti[i].value) if i < len(ti) else ""
                    t += ("：" if t else "") + str(c.value)
                    fields.append(t)
                if not fields:
                    continue
                line = "; ".join(fields)
                if sheetname.lower().find("sheet") < 0:
                    line += " ——" + sheetname
                res.append(line)
        return res

    @staticmethod
    def row_number(fnm, binary):
        if fnm.split(".")[-1].lower().find("xls") >= 0:
            wb = RAGFlowExcelParser._load_excel_to_workbook(BytesIO(binary))
            total = 0

            for sheetname in wb.sheetnames:
                try:
                    ws = wb[sheetname]
                    total += RAGFlowExcelParser._get_actual_row_count(ws)
                except Exception as e:
                    logging.warning(f"Skip sheet '{sheetname}' due to rows access error: {e}")
                    continue
            return total

        if fnm.split(".")[-1].lower() in ["csv", "txt"]:
            encoding = find_codec(binary)
            txt = binary.decode(encoding, errors="ignore")
            return len(txt.split("\n"))


if __name__ == "__main__":
    psr = RAGFlowExcelParser()
    psr(sys.argv[1])
