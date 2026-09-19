"""Successor integration: layout evidence must survive the real lifecycle."""
from io import BytesIO
from copy import deepcopy
import hashlib
import json
import re
from pathlib import Path
from datetime import datetime

import pytest

from openpyxl import Workbook

from rag.app import tabular_structure as producer
from rag.app.tabular_structure_runtime import _merge_sheet_projections
from test.fuxi.test_table_semantic_rows import _load_table_module
from test.fuxi.test_tabular_structure_projection import _VerifiedStorage


def _enable_candidate(monkeypatch):
    from rag.app.source_layout import SOURCE_LAYOUT_PROJECTION_CONTRACT
    for key, value in zip((
        "PRODUCER_SCHEMA_VERSION", "PROJECTION_VERSION",
        "STRUCTURE_PRODUCER_ALGORITHM_VERSION", "ENUMERATION_RULE_VERSION",
    ), SOURCE_LAYOUT_PROJECTION_CONTRACT):
        monkeypatch.setattr(producer, key, value)
    monkeypatch.setattr(producer, "_CURRENT_PROJECTION_CONTRACT", SOURCE_LAYOUT_PROJECTION_CONTRACT)


def _enable_reviewed_predecessor(monkeypatch):
    contract = ("table-producer/v6", "tabular-structure-projection/v6",
                "region-producer/v28", "enumeration-rules/v9")
    for key, value in zip((
        "PRODUCER_SCHEMA_VERSION", "PROJECTION_VERSION",
        "STRUCTURE_PRODUCER_ALGORITHM_VERSION", "ENUMERATION_RULE_VERSION",
    ), contract):
        monkeypatch.setattr(producer, key, value)
    monkeypatch.setattr(producer, "PROJECTION_PART_VERSION", "tabular-structure-part/v3")
    monkeypatch.setattr(producer, "_CURRENT_PROJECTION_CONTRACT", contract)


def _enable_named_candidate(monkeypatch):
    from rag.app.source_layout import SOURCE_LAYOUT_TITLE_PROJECTION_CONTRACT
    for key, value in zip(("PRODUCER_SCHEMA_VERSION", "PROJECTION_VERSION",
                           "STRUCTURE_PRODUCER_ALGORITHM_VERSION", "ENUMERATION_RULE_VERSION"),
                          SOURCE_LAYOUT_TITLE_PROJECTION_CONTRACT):
        monkeypatch.setattr(producer, key, value)
    monkeypatch.setattr(producer, "_CURRENT_PROJECTION_CONTRACT", SOURCE_LAYOUT_TITLE_PROJECTION_CONTRACT)


def test_merge_membership_is_checked_once_per_span():
    sheet = Workbook().active
    sheet["A1"] = "source declaration"
    sheet.merge_cells("A1:J10")

    class CountedMembers(set):
        checks = 0

        def __contains__(self, value):
            self.checks += 1
            return super().__contains__(value)

    members = CountedMembers((r, c) for r in range(1, 11) for c in range(1, 11))
    cells = producer._source_layout_cells(sheet, members)
    assert cells == [{"anchor": [1, 1], "span": [1, 1, 10, 10],
                      "value": "source declaration", "state": "literal"}]
    assert members.checks == len(members)


def test_released_default_emits_source_layout_and_keeps_predecessor_readable(monkeypatch):
    book = Workbook()
    book.active["A1"] = "Applicant"
    book.active["B1"] = "Alpha"
    stream = BytesIO()
    book.save(stream)
    projection = producer.build_tabular_structure_projection(
        "anonymous.xlsx", stream.getvalue(), parser=_load_table_module(monkeypatch).Excel(),
    )
    assert projection.get("source_layouts"), "released default must emit original form context"
    assert projection["structure_algorithm_version"] == "region-producer/v30"
    assert {layout["version"] for layout in projection["source_layouts"]} == {"source-layout/v2"}
    assert {layout["sheet_name"] for layout in projection["source_layouts"]} == {book.active.title}
    assert producer._RETAINED_SERVING_PROJECTION_CONTRACT == (
        "table-producer/v7", "tabular-structure-projection/v7", "region-producer/v29", "enumeration-rules/v9",
    )
    assert producer.PROJECTION_PART_VERSION == "tabular-structure-part/v4"
    for version in ("region-producer/v27", "region-producer/v28"):
        assert ("table-producer/v6", "tabular-structure-projection/v6", version,
                "enumeration-rules/v9") in producer._KNOWN_BACKFILL_PROJECTION_CONTRACTS


def test_membership_closure_expands_shared_merge_only_once(monkeypatch):
    sheet = Workbook().active
    sheet.merge_cells("A1:J10")
    yielded = 0
    def counted_range(*args):
        nonlocal yielded
        for value in range(*args):
            yielded += 1
            yield value
    monkeypatch.setattr(producer, "range", counted_range, raising=False)
    groups = producer._source_layout_memberships(sheet, [
        {"members": {(1, 1)}}, {"members": {(10, 10)}},
    ])
    assert groups == [{(r, c) for r in range(1, 11) for c in range(1, 11)}]
    assert yielded == 110  # Ten rows and their hundred coordinates, once.


def test_membership_closure_preserves_merge_bridges_and_disconnected_groups():
    sheet = Workbook().active
    sheet.merge_cells("A1:B2")
    sheet.merge_cells("D4:E5")
    regions = [{"members": {(1, 1)}}, {"members": {(2, 2), (4, 4)}},
               {"members": {(5, 5)}}, {"members": {(9, 9)}}]
    expected = [{(r, c) for r in (1, 2) for c in (1, 2)} |
                {(r, c) for r in (4, 5) for c in (4, 5)}, {(9, 9)}]
    assert producer._source_layout_memberships(sheet, regions) == expected
    assert producer._source_layout_memberships(sheet, list(reversed(regions))) == expected


def test_incomplete_merge_rejected_before_scanning_rectangle():
    sheet = Workbook().active
    sheet["A1"] = "source"
    from openpyxl.worksheet.cell_range import CellRange
    sheet.merged_cells.add(CellRange("A1:XFD1048576"))

    class NoRectangleScan(set):
        def __contains__(self, value):
            raise AssertionError("unbounded rectangle scan")

    with pytest.raises(ValueError, match="incomplete merge membership"):
        producer._source_layout_cells(sheet, NoRectangleScan({(1, 1)}))


def _layout_checkpoint(ordinal):
    generation = "bd0de022-ef46-4d59-8b1d-09c3761661cb"
    return {
        "version": "tabular-structure-projection/v7",
        "producer_schema_version": "table-producer/v7",
        "producer_generation_ref": generation,
        "structure_algorithm_version": "region-producer/v29",
        "enumeration_rule_version": "enumeration-rules/v9",
        "source_sha256": "a" * 64,
        "tables": [], "rows": [],
        "source_layouts": [{
            "version": "source-layout/v1", "source_sha256": "a" * 64,
            "producer_generation_ref": generation, "sheet_ordinal": ordinal,
            "object_ref": "layout_" + producer._versioned_digest(
                "source-layout-object/v1", "a" * 64, ordinal,
                producer._region_membership_sha256(ordinal, {(1, 1)}),
            ),
            "cells": [{"anchor": [1, 1], "span": [1, 1, 1, 1], "value": f"value-{ordinal}", "state": "literal"}],
        }],
    }


def test_checkpoint_merge_preserves_every_sheet_layout():
    first, second = _layout_checkpoint(1), _layout_checkpoint(2)
    merged = _merge_sheet_projections([first, second], generation_ref=first["producer_generation_ref"])
    assert merged["source_layouts"] == first["source_layouts"] + second["source_layouts"]


def test_named_layout_lifecycle_preserves_titles_and_retained_v29():
    from rag.app.source_layout import SOURCE_LAYOUT_TITLE_PROJECTION_CONTRACT
    first, second = _layout_checkpoint(1), _layout_checkpoint(2)
    for projection, title in ((first, "Declaration"), (second, "Declaration (variant)")):
        for key, value in zip(("producer_schema_version", "version",
                               "structure_algorithm_version", "enumeration_rule_version"),
                              SOURCE_LAYOUT_TITLE_PROJECTION_CONTRACT):
            projection[key] = value
        projection["source_layouts"][0].update(version="source-layout/v2", sheet_name=title)
    projection = _merge_sheet_projections([first, second], generation_ref=first["producer_generation_ref"])
    storage = _VerifiedStorage()
    receipt = producer.store_tabular_structure_projection(
        storage, bucket="bucket", document_id="document", projection=projection, rows_per_part=1,
    )
    restored = producer._load_tabular_structure_projection_for_contracts(
        storage, bucket="bucket", document_id="document",
        producer_generation_ref=first["producer_generation_ref"],
        manifest_object_name=receipt["manifest_object_name"],
        manifest_sha256=receipt["manifest_sha256"],
        accepted_contracts=frozenset({SOURCE_LAYOUT_TITLE_PROJECTION_CONTRACT}),
    )
    assert restored == projection
    inventory = producer.list_tabular_structure_projection_objects(
        storage, bucket="bucket", document_id="document",
        producer_generation_ref=first["producer_generation_ref"],
        manifest_object_name=receipt["manifest_object_name"],
        manifest_sha256=receipt["manifest_sha256"], expected_part_count=2,
    )
    assert len(inventory["object_names"]) == 3


@pytest.mark.parametrize("boundary", ["projection", "checkpoint"])
def test_named_layout_rejects_conflicting_titles_on_one_sheet(boundary):
    from rag.app.source_layout import SOURCE_LAYOUT_TITLE_PROJECTION_CONTRACT
    first = _layout_checkpoint(1)
    for key, value in zip(("producer_schema_version", "version",
                           "structure_algorithm_version", "enumeration_rule_version"),
                          SOURCE_LAYOUT_TITLE_PROJECTION_CONTRACT):
        first[key] = value
    first["source_layouts"][0].update(version="source-layout/v2", sheet_name="Declaration")
    second = deepcopy(first)
    layout = second["source_layouts"][0]
    layout["sheet_name"] = "Different title"
    layout["cells"][0].update(anchor=[3, 1], span=[3, 1, 3, 1])
    layout["object_ref"] = "layout_" + producer._versioned_digest(
        "source-layout-object/v1", "a" * 64, 1,
        producer._region_membership_sha256(1, {(3, 1)}),
    )
    with pytest.raises((ValueError, RuntimeError), match="title"):
        if boundary == "checkpoint":
            _merge_sheet_projections([first, second], generation_ref=first["producer_generation_ref"])
        else:
            first["source_layouts"].extend(second["source_layouts"])
            producer._validate_tabular_structure_projection_for_contract(first, SOURCE_LAYOUT_TITLE_PROJECTION_CONTRACT)


@pytest.mark.parametrize("title", ["Declaration", "申报表（修订版）", "Independent source"])
def test_successor_builder_preserves_original_sheet_title(monkeypatch, title):
    from rag.app.source_layout import SOURCE_LAYOUT_TITLE_PROJECTION_CONTRACT
    for key, value in zip(("PRODUCER_SCHEMA_VERSION", "PROJECTION_VERSION",
                           "STRUCTURE_PRODUCER_ALGORITHM_VERSION", "ENUMERATION_RULE_VERSION"),
                          SOURCE_LAYOUT_TITLE_PROJECTION_CONTRACT):
        monkeypatch.setattr(producer, key, value)
    monkeypatch.setattr(producer, "_CURRENT_PROJECTION_CONTRACT", SOURCE_LAYOUT_TITLE_PROJECTION_CONTRACT)
    book = Workbook()
    book.active.title = title
    book.active["A1"] = "Applicant"
    book.active["B1"] = "Alpha"
    stream = BytesIO()
    book.save(stream)
    projection = producer.build_tabular_structure_projection(
        "anonymous.xlsx", stream.getvalue(), parser=_load_table_module(monkeypatch).Excel(),
    )
    assert projection.get("source_layouts")
    assert {layout["sheet_name"] for layout in projection["source_layouts"]} == {title}
    assert {layout["version"] for layout in projection["source_layouts"]} == {"source-layout/v2"}


@pytest.mark.parametrize("damage", ["identity", "moved", "missing"])
def test_candidate_projection_recomputes_layout_membership_identity(damage):
    from rag.app.source_layout import SOURCE_LAYOUT_PROJECTION_CONTRACT
    projection = _layout_checkpoint(1)
    layout = projection["source_layouts"][0]
    if damage == "identity":
        layout["object_ref"] = "layout_" + "0" * 64
    elif damage == "moved":
        layout["cells"][0].update(anchor=[2, 1], span=[2, 1, 2, 1])
    else:
        layout["cells"] = []
    with pytest.raises(ValueError, match="layout.*identity"):
        producer._validate_tabular_structure_projection_for_contract(projection, SOURCE_LAYOUT_PROJECTION_CONTRACT)


def test_layout_digest_matches_source_membership_with_interleaved_merges_and_details():
    spans = [[1, 1, 3, 1], [1, 3, 2, 4], [1, 1, 3, 1], [7, 2, 7, 2]]
    members = {(row, column) for top, left, bottom, right in spans
               for row in range(top, bottom + 1) for column in range(left, right + 1)}
    layout = {"sheet_ordinal": 3, "cells": [{"span": span} for span in spans]}
    assert producer._source_layout_membership_sha256(layout) == producer._region_membership_sha256(3, members)


def test_candidate_deletion_inventory_survives_part_already_removed():
    projection = _layout_checkpoint(1)
    storage = _VerifiedStorage()
    receipt = producer.store_tabular_structure_projection(
        storage, bucket="bucket", document_id="document", projection=projection,
    )
    part_key = next(key for key in storage.objects if "/part-" in key[1])
    del storage.objects[part_key]
    inventory = producer.list_tabular_structure_projection_objects(
        storage, bucket="bucket", document_id="document",
        producer_generation_ref=projection["producer_generation_ref"],
        manifest_object_name=receipt["manifest_object_name"],
        manifest_sha256=receipt["manifest_sha256"], expected_part_count=1,
    )
    assert inventory["object_names"] == [part_key[1], receipt["manifest_object_name"]]


@pytest.mark.parametrize("damage", ["scope", "offset", "count", "source", "generation", "duplicate", "version"])
def test_candidate_deletion_inventory_rejects_rehashed_manifest_drift(damage):
    projection = _layout_checkpoint(1)
    storage = _VerifiedStorage()
    receipt = producer.store_tabular_structure_projection(
        storage, bucket="bucket", document_id="document", projection=projection,
    )
    manifest = json.loads(storage.objects[("bucket", receipt["manifest_object_name"])])
    if damage == "scope":
        manifest["parts"][0]["object_name"] = "another-document/part.json"
    elif damage == "offset":
        manifest["parts"][0]["layout_offset"] = 1
    elif damage == "count":
        manifest["source_layouts"][0]["cell_count"] = 2
    elif damage == "source":
        manifest["source_layouts"][0]["source_sha256"] = "b" * 64
    elif damage == "generation":
        manifest["source_layouts"][0]["producer_generation_ref"] = "another"
    elif damage == "duplicate":
        manifest["source_layouts"].append(deepcopy(manifest["source_layouts"][0]))
    else:
        manifest["structure_algorithm_version"] = "region-producer/v28"
    payload = json.dumps(manifest).encode()
    digest = hashlib.sha256(payload).hexdigest()
    name = receipt["manifest_object_name"].replace(receipt["manifest_sha256"], digest)
    storage.objects[("bucket", name)] = payload
    with pytest.raises(producer.StructureSnapshotChanged):
        producer.list_tabular_structure_projection_objects(
            storage, bucket="bucket", document_id="document",
            producer_generation_ref=projection["producer_generation_ref"],
            manifest_object_name=name, manifest_sha256=digest,
        )


def test_candidate_layout_only_generation_roundtrips_without_record_counts():
    from rag.app.source_layout import SOURCE_LAYOUT_PROJECTION_CONTRACT
    first, second = _layout_checkpoint(1), _layout_checkpoint(2)
    projection = _merge_sheet_projections([first, second], generation_ref=first["producer_generation_ref"])
    storage = _VerifiedStorage()
    receipt = producer.store_tabular_structure_projection(
        storage, bucket="bucket", document_id="document", projection=projection, rows_per_part=1,
    )
    restored = producer._load_tabular_structure_projection_for_contracts(
        storage, bucket="bucket", document_id="document",
        producer_generation_ref=first["producer_generation_ref"],
        manifest_object_name=receipt["manifest_object_name"],
        manifest_sha256=receipt["manifest_sha256"], expected_part_count=receipt["part_count"],
        accepted_contracts=frozenset({SOURCE_LAYOUT_PROJECTION_CONTRACT}),
    )
    assert restored == projection
    assert receipt["row_count"] == 0
    assert receipt["part_count"] == 2


@pytest.mark.parametrize("damage", ["missing", "corrupt"])
def test_candidate_layout_part_damage_is_rejected(damage):
    from rag.app.source_layout import SOURCE_LAYOUT_PROJECTION_CONTRACT
    projection = _layout_checkpoint(1)
    storage = _VerifiedStorage()
    receipt = producer.store_tabular_structure_projection(
        storage, bucket="bucket", document_id="document", projection=projection,
    )
    part_key = next(key for key in storage.objects if "/part-" in key[1])
    if damage == "missing":
        del storage.objects[part_key]
    else:
        storage.objects[part_key] = b"{}"
    with pytest.raises((producer.StructureSnapshotMissing, producer.StructureSnapshotChanged)):
        producer._load_tabular_structure_projection_for_contracts(
            storage, bucket="bucket", document_id="document",
            producer_generation_ref=projection["producer_generation_ref"],
            manifest_object_name=receipt["manifest_object_name"],
            manifest_sha256=receipt["manifest_sha256"],
            accepted_contracts=frozenset({SOURCE_LAYOUT_PROJECTION_CONTRACT}),
        )


def test_readback_rejects_moved_layout_even_when_all_storage_hashes_are_recomputed():
    from rag.app.source_layout import SOURCE_LAYOUT_PROJECTION_CONTRACT
    projection = _layout_checkpoint(1)
    storage = _VerifiedStorage()
    receipt = producer.store_tabular_structure_projection(
        storage, bucket="bucket", document_id="document", projection=projection,
    )
    manifest = json.loads(storage.objects[("bucket", receipt["manifest_object_name"])])
    entry = manifest["parts"][0]
    part = json.loads(storage.objects[("bucket", entry["object_name"])])
    part["layout_cells"][0]["cell"].update(anchor=[2, 1], span=[2, 1, 2, 1])
    payload = json.dumps(part).encode()
    digest = hashlib.sha256(payload).hexdigest()
    entry["object_name"] = entry["object_name"].replace(entry["sha256"], digest)
    entry["sha256"] = digest
    storage.objects[("bucket", entry["object_name"])] = payload
    payload = json.dumps(manifest).encode()
    digest = hashlib.sha256(payload).hexdigest()
    name = receipt["manifest_object_name"].replace(receipt["manifest_sha256"], digest)
    storage.objects[("bucket", name)] = payload
    with pytest.raises(producer.StructureSnapshotChanged, match="validation changed") as failure:
        producer._load_tabular_structure_projection_for_contracts(
            storage, bucket="bucket", document_id="document",
            producer_generation_ref=projection["producer_generation_ref"],
            manifest_object_name=name, manifest_sha256=digest,
            accepted_contracts=frozenset({SOURCE_LAYOUT_PROJECTION_CONTRACT}),
        )
    assert str(failure.value.__cause__) == "source layout membership identity mismatch"


@pytest.mark.parametrize("drift", ["missing", "generation", "source", "duplicate", "order"])
def test_checkpoint_merge_rejects_layout_drift(drift):
    first, second = _layout_checkpoint(1), _layout_checkpoint(2)
    if drift == "missing":
        del second["source_layouts"]
    elif drift == "generation":
        second["source_layouts"][0]["producer_generation_ref"] = "another"
    elif drift == "source":
        second["source_layouts"][0]["source_sha256"] = "b" * 64
    elif drift == "duplicate":
        second["source_layouts"] = deepcopy(first["source_layouts"])
    elif drift == "order":
        first, second = second, first
    with pytest.raises((ValueError, RuntimeError)):
        _merge_sheet_projections([first, second], generation_ref=first["producer_generation_ref"])


def test_layout_survives_sheet_checkpoint_merge_and_immutable_readback(monkeypatch):
    # Exercise the successor before capability rotation, never add new fields
    # to a manifest advertising the deployed v28 contract.
    _enable_candidate(monkeypatch)
    book = Workbook()
    for sheet, label, detail in (
        (book.active, "Applicant", "Alpha"),
        (book.create_sheet(), "Reviewer", "Beta"),
    ):
        sheet.merge_cells("A1:C1")
        sheet["A1"] = label
        sheet.merge_cells("A2:C2")
        sheet["A2"] = detail
        sheet["A4"] = "Date"
        sheet["B4"] = "/"
    stream = BytesIO()
    book.save(stream)
    binary = stream.getvalue()
    parser = _load_table_module(monkeypatch).Excel()
    generation = "bd0de022-ef46-4d59-8b1d-09c3761661cb"
    checkpoints = [producer.build_tabular_structure_projection(
        "anonymous.xlsx", binary, parser=parser,
        producer_generation_ref=generation, sheet_ordinals={ordinal},
    ) for ordinal in (1, 2)]
    for ordinal, checkpoint in enumerate(checkpoints, 1):
        assert checkpoint.get("source_layouts"), "Producer omitted source layout evidence"
        assert {layout["sheet_ordinal"] for layout in checkpoint["source_layouts"]} == {ordinal}
    merged = _merge_sheet_projections(checkpoints, generation_ref=generation)
    assert merged["source_layouts"] == checkpoints[0]["source_layouts"] + checkpoints[1]["source_layouts"]
    storage = _VerifiedStorage()
    receipt = producer.store_tabular_structure_projection(
        storage, bucket="bucket", document_id="document", projection=merged, rows_per_part=1,
    )
    restored = producer.load_tabular_structure_projection(
        storage, bucket="bucket", document_id="document",
        producer_generation_ref=generation,
        manifest_object_name=receipt["manifest_object_name"],
        manifest_sha256=receipt["manifest_sha256"],
        expected_part_count=receipt["part_count"],
    )
    assert restored == merged
    values = [cell["value"] for layout in restored["source_layouts"] for cell in layout["cells"]]
    assert "Alpha" in values and "Beta" in values


@pytest.mark.parametrize("ordinal", [6, 7])
def test_current_pn01_layout_build_storage_preserves_native_values(monkeypatch, ordinal, tmp_path):
    from python_calamine import CalamineWorkbook
    source = Path("/opt/fuxi/evidence/g91-source-download-20260916/G91-PN01.xls")
    if not source.exists():
        pytest.skip("exact uploaded source unavailable")
    binary = source.read_bytes()
    assert hashlib.sha256(binary).hexdigest() == "0f0bab58c1eb8c67541c3fae9e068f2fe274228c615ea95ff7ad9b2f99699cdc"
    _enable_candidate(monkeypatch)
    projection = producer.build_tabular_structure_projection(
        "anonymous.xls", binary, parser=_load_table_module(monkeypatch).Excel(),
        sheet_ordinals={ordinal},
    )
    storage = _VerifiedStorage()
    receipt = producer.store_tabular_structure_projection(
        storage, bucket="bucket", document_id="document", projection=projection, rows_per_part=13,
    )
    restored = producer.load_tabular_structure_projection(
        storage, bucket="bucket", document_id="document",
        producer_generation_ref=projection["producer_generation_ref"],
        manifest_object_name=receipt["manifest_object_name"], manifest_sha256=receipt["manifest_sha256"],
    )
    assert restored == projection
    # Verify actual Producer bytes through the separately maintained release
    # gate, not a second hand-written representation of the new protocol.
    import importlib.util
    verifier_path = Path("/opt/fuxi/offline-delivery/codex-skills/fuxi-tabular-knowledge-qa-closure/scripts/verify_shadow_generation_closure.py")
    assert verifier_path.is_file(), "release verifier required for source closure integration"
    spec = importlib.util.spec_from_file_location("shadow_release_gate", verifier_path)
    verifier = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(verifier)
    manifest_bytes = storage.get("bucket", receipt["manifest_object_name"])
    manifest = json.loads(manifest_bytes)
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_bytes(manifest_bytes)
    part_paths = []
    for index, part in enumerate(manifest["parts"]):
        part_path = tmp_path / f"part-{index}.json"
        part_path.write_bytes(storage.get("bucket", part["object_name"]))
        part_paths.append(part_path)
    digest, _, contract, summary = verifier.validate_source_closure(manifest_path, part_paths)
    assert digest == receipt["manifest_sha256"]
    assert contract["projection_part"] == "tabular-structure-part/v4"
    assert summary["part_count"] == receipt["part_count"]
    cells = [cell for layout in restored["source_layouts"] for cell in layout["cells"]]
    worksheet = _load_table_module(monkeypatch).Excel()._load_excel_to_workbook(BytesIO(binary)).worksheets[ordinal - 1]
    blank_borders = {(cell.row, cell.column) for cell in worksheet._cells.values()
                     if cell.value is None and any(
                         getattr(getattr(cell.border, side, None), "style", None)
                         for side in ("left", "right", "top", "bottom"))}
    assert blank_borders
    for row, column in blank_borders:
        assert any(cell["span"][0] <= row <= cell["span"][2]
                   and cell["span"][1] <= column <= cell["span"][3] for cell in cells), (ordinal, row, column)
    rows = CalamineWorkbook.from_filelike(BytesIO(binary)).get_sheet_by_index(ordinal - 1).to_python(skip_empty_area=False)
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row, 1):
            if value is None or not str(value).strip():
                continue
            assert any(cell["span"][0] <= r <= cell["span"][2]
                       and cell["span"][1] <= c <= cell["span"][3]
                       and str(cell["value"]) == str(value) for cell in cells), (ordinal, r, c)


@pytest.mark.parametrize("named", [False, True])
@pytest.mark.parametrize("merged", [False, True])
def test_builder_preserves_uncached_formula_as_unresolved(monkeypatch, merged, named):
    (_enable_named_candidate if named else _enable_candidate)(monkeypatch)
    book = Workbook()
    book.active["A1"] = "Calculated value"
    book.active["B1"] = "=1+1"
    if merged:
        book.active.merge_cells("B1:C2")
    stream = BytesIO()
    book.save(stream)
    projection = producer.build_tabular_structure_projection(
        "anonymous.xlsx", stream.getvalue(), parser=_load_table_module(monkeypatch).Excel(),
    )
    cells = [cell for layout in projection["source_layouts"] for cell in layout["cells"]]
    formula = next(cell for cell in cells if cell["anchor"] == [1, 2])
    assert formula["state"] == "formula_unresolved"
    assert formula["value"] == "=1+1"


def test_builder_date_value_survives_immutable_storage(monkeypatch):
    _enable_candidate(monkeypatch)
    book = Workbook()
    book.active["A1"] = "Date"
    book.active["B1"] = datetime(2026, 9, 18, 12, 34, 56)
    stream = BytesIO()
    book.save(stream)
    projection = producer.build_tabular_structure_projection(
        "anonymous.xlsx", stream.getvalue(), parser=_load_table_module(monkeypatch).Excel(),
    )
    storage = _VerifiedStorage()
    receipt = producer.store_tabular_structure_projection(
        storage, bucket="bucket", document_id="document", projection=projection,
    )
    restored = producer.load_tabular_structure_projection(
        storage, bucket="bucket", document_id="document",
        producer_generation_ref=projection["producer_generation_ref"],
        manifest_object_name=receipt["manifest_object_name"], manifest_sha256=receipt["manifest_sha256"],
    )
    cells = [cell for layout in restored["source_layouts"] for cell in layout["cells"]]
    cell = next(cell for cell in cells if cell["anchor"] == [1, 2])
    assert cell["value"] == "2026-09-18T12:34:56"
    assert cell["state"] == "datetime"


def test_current_source_layout_regions_have_complete_merge_ownership(monkeypatch):
    source = Path("/opt/fuxi/evidence/g91-source-download-20260916/G91-PN01.xls")
    if not source.exists():
        pytest.skip("exact uploaded source unavailable; not acceptance evidence")
    binary = source.read_bytes()
    assert hashlib.sha256(binary).hexdigest() == "0f0bab58c1eb8c67541c3fae9e068f2fe274228c615ea95ff7ad9b2f99699cdc"
    parser = _load_table_module(monkeypatch).Excel()
    _enable_candidate(monkeypatch)
    book = parser._load_excel_to_workbook(BytesIO(binary))
    conflicts = []
    for ordinal, sheet in enumerate(book.worksheets, start=1):
        regions = producer._worksheet_structure_regions(parser, sheet, ordinal)
        for members in producer._source_layout_memberships(sheet, regions):
            for merge in sheet.merged_cells.ranges:
                span = {(r, c) for r in range(merge.min_row, merge.max_row + 1)
                        for c in range(merge.min_col, merge.max_col + 1)}
                if span & members and not span <= members:
                    conflicts.append((ordinal, str(merge), len(span & members), len(span),
                                      sheet.cell(merge.min_row, merge.min_col).value is None))
    assert not conflicts, conflicts


@pytest.mark.parametrize("source_path,expected_sha", [
    ("/opt/fuxi/evidence/ppap-four-file-current-identity-20260914t051021853198z-6d7431ba/remote-originals/D91-轮胎总成PH01-中策橡胶.xls", "4efb1838b34a777e6d39f60f24c7e97d4873f4c726203cb6f418e5409762649b"),
    ("/opt/fuxi/evidence/ppap-four-file-current-identity-20260914t051021853198z-6d7431ba/remote-originals/D91-前稳定杆接头总成PH01-江西荣成.xls", "365251db57fb72a03219625d21879bf06e656f42f848a34372d3dd1a889d0d81"),
    ("/opt/fuxi/evidence/f515-new-source-producer-replay-20260915-20260915t020918747155z-f573fe62/F515-new.xls", "36a044c8dd2ef8f1c134eb46819477c3eca222f2f2c0bea8f18804bac22d9d37"),
    ("/opt/fuxi/evidence/g91-source-download-20260916/G91-PN01.xls", "0f0bab58c1eb8c67541c3fae9e068f2fe274228c615ea95ff7ad9b2f99699cdc"),
], ids=["tire", "front", "replacement", "column"])
def test_successor_preserves_every_existing_list_field_on_reviewed_sources(monkeypatch, source_path, expected_sha):
    source = Path(source_path)
    if not source.exists():
        pytest.skip("exact uploaded source unavailable; not acceptance evidence")
    binary = source.read_bytes()
    assert hashlib.sha256(binary).hexdigest() == expected_sha
    parser = _load_table_module(monkeypatch).Excel()
    generation = "11111111-1111-5111-8111-111111111111"
    _enable_candidate(monkeypatch)
    print(f"source-replay {expected_sha} legacy-start", flush=True)
    previous = producer.build_tabular_structure_projection("anonymous.xls", binary, parser=parser, producer_generation_ref=generation)
    assert previous["structure_algorithm_version"] == "region-producer/v29"
    assert previous["source_layouts"]
    print(f"source-replay {expected_sha} legacy-complete tables={len(previous['tables'])} rows={len(previous['rows'])}", flush=True)
    _enable_named_candidate(monkeypatch)
    print(f"source-replay {expected_sha} successor-start", flush=True)
    successor = producer.build_tabular_structure_projection("anonymous.xls", binary, parser=parser, producer_generation_ref=generation)
    assert successor["structure_algorithm_version"] == "region-producer/v30"
    print(f"source-replay {expected_sha} successor-complete tables={len(successor['tables'])} rows={len(successor['rows'])}", flush=True)

    def list_projection(projection):
        # Version-bound IDs rotate by contract. Preserve membership digests,
        # all source fields, counts, labels, columns and grouped parent ordinals.
        identities = {}
        for table in projection["tables"]:
            membership = re.fullmatch(r"tbl_v2_([0-9a-f]{64})_[0-9a-f]{64}", table["table_ref"])
            assert membership is not None
            digest = membership.group(1)
            assert table["table_ref"] == producer._table_ref_for_contract(
                projection["source_sha256"], table["sheet_ordinal"], table["table_ordinal"], digest,
                producer_schema_version=projection["producer_schema_version"],
                projection_version=projection["version"],
                structure_algorithm_version=projection["structure_algorithm_version"],
                enumeration_rule_version=projection["enumeration_rule_version"],
            )
            identities[table["table_ref"]] = (
                f"sheet:{table['sheet_ordinal']}:table:{table['table_ordinal']}:members:{digest}"
            )
        for record in projection["rows"]:
            assert record["id"] == "tsr_v1_" + producer._versioned_digest(
                "tabular-row-record/v1", generation, record["row_ref_kwd"])
        def normalize(value, key=None):
            if isinstance(value, dict):
                return {field: normalize(value["row_ref_kwd"] if field == "id" and "row_ref_kwd" in value else item, field) for field, item in value.items()
                        if field not in {"producer_schema_version_kwd", "projection_version_kwd", "structure_algorithm_version_kwd", "enumeration_rule_version_kwd"}}
            if isinstance(value, list):
                return [normalize(item) for item in value]
            if isinstance(value, str):
                for identity, replacement in identities.items():
                    if value == identity or value.startswith(identity + ":"):
                        return value.replace(identity, replacement, 1)
            return value
        return normalize({"tables": projection["tables"], "rows": projection["rows"]})

    assert list_projection(successor) == list_projection(previous)
    assert successor["source_layouts"]
