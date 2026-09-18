"""Source geometry is evidence, not permission to infer field semantics."""
import hashlib
from io import BytesIO
from pathlib import Path
import pytest
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Border, Side, Font

from rag.app import tabular_structure
from rag.app.source_layout import validate_source_layout, validate_layout_membership
from test.fuxi.test_table_semantic_rows import _load_table_module


def test_form_cells_preserve_merge_blank_and_duplicate_text_without_outside_values():
    sheet = Workbook().active
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Date"
    sheet["A2"] = " 2026-09-18 "
    sheet["C1"] = "Date"
    sheet["C2"] = "/"
    sheet["D1"] = "outside"
    members = {(1, 1), (1, 2), (2, 1), (2, 2), (1, 3), (2, 3)}
    cells = tabular_structure._source_layout_cells(sheet, members)
    assert cells == [
        {"anchor": [1, 1], "span": [1, 1, 1, 2], "value": "Date", "state": "literal"},
        {"anchor": [1, 3], "span": [1, 3, 1, 3], "value": "Date", "state": "literal"},
        {"anchor": [2, 1], "span": [2, 1, 2, 1], "value": " 2026-09-18 ", "state": "literal"},
        {"anchor": [2, 2], "span": [2, 2, 2, 2], "value": None, "state": "blank"},
        {"anchor": [2, 3], "span": [2, 3, 2, 3], "value": "/", "state": "literal"},
    ]


def test_form_cells_reject_partial_merge_membership():
    sheet = Workbook().active
    sheet.merge_cells("A1:B2")
    sheet["A1"] = "Declaration"
    with pytest.raises(ValueError, match="merge membership"):
        tabular_structure._source_layout_cells(sheet, {(1, 1), (1, 2)})


def test_layout_membership_closes_blank_anchor_merge_without_changing_list_regions():
    sheet = Workbook().active
    sheet.merge_cells("A1:A2")
    sheet._cells[(2, 1)] = Cell(sheet, row=2, column=1, value="Source detail")
    regions = [{"members": {(2, 1)}}, {"members": {(8, 8)}}]
    memberships = tabular_structure._source_layout_memberships(sheet, regions)
    assert memberships == [{(1, 1), (2, 1)}, {(8, 8)}]
    assert regions == [{"members": {(2, 1)}}, {"members": {(8, 8)}}]
    cells = tabular_structure._source_layout_cells(sheet, memberships[0])
    assert [cell["value"] for cell in cells] == [None, "Source detail"]


def test_layout_preserves_bordered_blank_slots_without_promoting_styles_to_values():
    sheet = Workbook().active
    sheet["A1"] = "Declaration"
    sheet["B1"].border = Border(bottom=Side(style="thin"))
    sheet["H8"].border = Border(left=Side(style="thin"))
    sheet["Z99"].font = Font(bold=True)
    regions = [{"members": {(1, 1)}}]
    memberships = tabular_structure._source_layout_memberships(sheet, regions)
    assert set().union(*memberships) == {(1, 1), (1, 2), (8, 8)}
    assert regions == [{"members": {(1, 1)}}]
    cells = [cell for members in memberships for cell in tabular_structure._source_layout_cells(sheet, members)]
    assert next(cell for cell in cells if cell["anchor"] == [1, 2])["state"] == "blank"
    assert next(cell for cell in cells if cell["anchor"] == [8, 8])["value"] is None


def test_layout_membership_unifies_shared_merge_without_filling_bounding_rectangle():
    sheet = Workbook().active
    sheet.merge_cells("A1:A3")
    regions = [{"members": {(1, 1), (1, 3)}}, {"members": {(3, 1)}}, {"members": {(9, 9)}}]
    assert tabular_structure._source_layout_memberships(sheet, regions) == [
        {(1, 1), (2, 1), (3, 1), (1, 3)}, {(9, 9)},
    ]


def test_form_cells_preserve_distinct_physical_values_inside_legacy_merge():
    sheet = Workbook().active
    sheet.merge_cells("A1:A2")
    sheet["A1"] = "Upper: 47.5"
    # The legacy loader can retain physical BIFF values inside merge ranges.
    sheet._cells[(2, 1)] = Cell(sheet, row=2, column=1, value="Lower: 45")
    cells = tabular_structure._source_layout_cells(sheet, {(1, 1), (2, 1)})
    assert [cell["value"] for cell in cells] == ["Upper: 47.5", "Lower: 45"]
    assert cells[1]["anchor"] == [2, 1]
    assert cells[1]["span"] == [1, 1, 2, 1]


def test_form_cells_do_not_treat_formula_text_as_cached_result():
    sheet = Workbook().active
    sheet["A1"] = "=1+1"
    sheet["B1"] = False
    cells = tabular_structure._source_layout_cells(sheet, {(1, 1), (1, 2)})
    assert cells[0]["state"] == "formula_unresolved"
    assert cells[0]["value"] == "=1+1"
    assert cells[1]["value"] is False


def test_unavailable_formula_expression_is_not_a_blank_field():
    sheet = Workbook().active
    cells = tabular_structure._source_layout_cells(
        sheet, {(1, 1)}, unresolved_formulas={(1, 1)},
    )
    assert cells == [{"anchor": [1, 1], "span": [1, 1, 1, 1],
                      "value": None, "state": "formula_unresolved"}]
    scope = {"source_sha256": "a" * 64, "producer_generation_ref": "generation",
             "sheet_ordinal": 1, "object_ref": "object"}
    assert validate_source_layout({"version": "source-layout/v1", **scope, "cells": cells}, **scope)["cells"] == cells


@pytest.mark.parametrize("sheet_ordinal", [6, 7])
def test_current_pn01_form_source_values_match_independent_reader(monkeypatch, sheet_ordinal):
    from python_calamine import CalamineWorkbook

    source = Path("/opt/fuxi/evidence/g91-source-download-20260916/G91-PN01.xls")
    if not source.exists():
        pytest.skip("exact uploaded source unavailable; not source acceptance")
    binary = source.read_bytes()
    assert hashlib.sha256(binary).hexdigest() == "0f0bab58c1eb8c67541c3fae9e068f2fe274228c615ea95ff7ad9b2f99699cdc"
    parser = _load_table_module(monkeypatch).Excel()
    worksheet = parser._load_excel_to_workbook(BytesIO(binary)).worksheets[sheet_ordinal - 1]
    members = tabular_structure._logical_occupied_cells(parser, worksheet)
    cells = tabular_structure._source_layout_cells(worksheet, members)
    scope = {
        "source_sha256": hashlib.sha256(binary).hexdigest(),
        "producer_generation_ref": "source-regression-generation",
        "sheet_ordinal": sheet_ordinal,
        "object_ref": "source-regression-object",
    }
    layout = validate_source_layout(
        {"version": "source-layout/v1", **scope, "cells": cells}, **scope,
    )
    validate_layout_membership(layout, members)
    cells = layout["cells"]
    independent = CalamineWorkbook.from_filelike(BytesIO(binary)).get_sheet_by_index(sheet_ordinal - 1)
    rows = independent.to_python(skip_empty_area=False)
    checked = 0
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            if value is None or not str(value).strip():
                continue
            matching = [cell for cell in cells if cell["span"][0] <= r <= cell["span"][2]
                        and cell["span"][1] <= c <= cell["span"][3]
                        and str(cell["value"]) == str(value)]
            assert matching, (sheet_ordinal, r, c)
            checked += 1
    assert checked > 0
